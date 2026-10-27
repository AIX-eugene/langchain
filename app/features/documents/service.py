"""Shared in-memory document extraction and Qdrant ingestion."""

import io
import logging
import threading
import uuid
from pathlib import PurePosixPath

from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter
from pypdf import PdfReader
from qdrant_client.models import Distance, FieldCondition, Filter, MatchValue, PointIdsList, PointStruct, VectorParams

from app.core.config import COLLECTION, MAX_DOCUMENT_MB
from app.core.config import MACEA_COLLECTION
from app.infrastructure.qdrant.store import make_client, make_embeddings

UPLOAD_COLLECTION = MACEA_COLLECTION
TYPES = {'.pdf', '.txt', '.md', '.docx', '.xlsx'}
BATCH_SIZE = 32
ingestion_lock = threading.Lock()
logger = logging.getLogger(__name__)


def source_name(filename):
    name = PurePosixPath((filename or '').replace('\\', '/')).name
    name = ''.join(c for c in name if c.isprintable()).strip()
    if not name or name in {'.', '..'} or len(name) > 255:
        raise ValueError('Use a filename between 1 and 255 characters.')
    if PurePosixPath(name).suffix.lower() not in TYPES:
        raise ValueError('Supported files: PDF, TXT, MD, DOCX, and XLSX.')
    return name


def read_upload(file):
    data = file.read(MAX_DOCUMENT_MB * 1024 * 1024 + 1)
    if len(data) > MAX_DOCUMENT_MB * 1024 * 1024:
        raise ValueError(f'The document is larger than {MAX_DOCUMENT_MB} MB.')
    if not data:
        raise ValueError('The document is empty.')
    return data


def extract_documents(data, name):
    suffix = PurePosixPath(name).suffix.lower()
    def doc(text, **metadata):
        return Document(page_content=text, metadata={'source': name, **metadata})
    try:
        if suffix == '.pdf':
            reader = PdfReader(io.BytesIO(data))
            if reader.is_encrypted and not reader.decrypt(''):
                raise ValueError('This PDF requires a password. Upload an unencrypted copy; password entry is not supported.')
            docs = [doc(page.extract_text() or '', page=i) for i, page in enumerate(reader.pages, 1)]
        elif suffix in {'.txt', '.md'}:
            docs = [doc(data.decode('utf-8-sig'))]
        elif suffix == '.docx':
            from docx import Document as WordDocument
            from docx.table import Table
            word = WordDocument(io.BytesIO(data))
            parts = []
            for block in word.iter_inner_content():
                parts.append('\n'.join('\t'.join(cell.text for cell in row.cells) for row in block.rows)
                             if isinstance(block, Table) else block.text)
            docs = [doc('\n\n'.join(parts))]
        elif suffix == '.xlsx':
            from openpyxl import load_workbook
            book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            try:
                docs = []
                for sheet in book:
                    for number, row in enumerate(sheet.iter_rows(values_only=True), 1):
                        values = [str(value) if value is not None else '' for value in row]
                        if any(value.strip() for value in values):
                            docs.append(doc(f'{sheet.title}, row {number}: ' + '\t'.join(values), sheet=sheet.title, row=number))
            finally:
                book.close()
        else:
            raise ValueError('Unsupported document type.')
    except UnicodeDecodeError as exc:
        raise ValueError('Text files must use UTF-8 encoding.') from exc
    except ImportError as exc:
        logger.exception('Document parser dependency unavailable for %s', suffix)
        raise ValueError(
            f'The {suffix} parser is unavailable in the running Python environment. '
            'Run uv sync, then restart the app with uv run server.py.'
        ) from exc
    except ValueError:
        raise
    except Exception as exc:
        logger.exception('Document extraction failed for %s', suffix)
        detail = ' '.join(str(exc).split())[:300] or 'No additional parser details.'
        raise ValueError(
            f'Could not read this {suffix} document: {type(exc).__name__}: {detail} '
            'Try opening it in its original application and saving a new copy.'
        ) from exc
    chunks = RecursiveCharacterTextSplitter(chunk_size=1000, chunk_overlap=200).split_documents(docs)
    chunks = [chunk for chunk in chunks if chunk.page_content.strip()]
    if not chunks:
        raise ValueError('No readable text found. Scanned PDFs need OCR first; spreadsheets need text or cached values.')
    return chunks


def ingest_document(data, filename, replace=False, client=None, embeddings=None, confirm_add=None, collection=COLLECTION):
    """Yield progress dictionaries. Consume fully: writes run independently of the browser."""
    name = source_name(filename)
    yield {'type': 'stage', 'message': 'Validating document'}
    # One writer per application process, including CLI callers using this helper.
    with ingestion_lock:
        client = client or make_client()
        embeddings = embeddings or make_embeddings()
        exists = client.collection_exists(collection)
        same = Filter(must=[FieldCondition(key='metadata.source', match=MatchValue(value=name))])
        old_ids = []
        if exists:
            offset = None
            while True:
                points, offset = client.scroll(collection, scroll_filter=same, offset=offset, limit=256,
                                               with_payload=False, with_vectors=False)
                old_ids.extend(point.id for point in points)
                if offset is None:
                    break
        if old_ids and not replace:
            yield {'type': 'duplicate', 'filename': name, 'passages': len(old_ids),
                   'message': f'{name} already exists. Confirm replacement to continue.'}
            return
        if exists and not old_ids and confirm_add and client.count(collection).count:
            if not confirm_add():
                raise ValueError('Cancelled.')
        yield {'type': 'stage', 'message': 'Extracting text'}
        chunks = extract_documents(data, name)
        vectors = []
        for start in range(0, len(chunks), BATCH_SIZE):
            batch = embeddings.embed_documents([c.page_content for c in chunks[start:start+BATCH_SIZE]])
            if len(batch) != len(chunks[start:start+BATCH_SIZE]) or any(not v for v in batch):
                raise ValueError('Embedding service returned incomplete vectors.')
            vectors.extend(batch)
            yield {'type': 'progress', 'stage': 'embedding', 'completed': len(vectors), 'total': len(chunks)}
        if not exists:
            client.create_collection(collection, vectors_config=VectorParams(size=len(vectors[0]), distance=Distance.COSINE))
        ids = [str(uuid.uuid4()) for _ in chunks]
        try:
            for start in range(0, len(chunks), BATCH_SIZE):
                end = min(start+BATCH_SIZE, len(chunks))
                client.upsert(collection, points=[PointStruct(id=ids[i], vector=vectors[i],
                              payload={'content': chunks[i].page_content, 'metadata': chunks[i].metadata})
                              for i in range(start, end)], wait=True)
                yield {'type': 'progress', 'stage': 'storage', 'completed': end, 'total': len(chunks)}
        except Exception as exc:
            try:
                client.delete(collection, points_selector=PointIdsList(points=ids), wait=True)
            except Exception as cleanup:
                raise RuntimeError(f'Upload failed and cleanup is incomplete. New point IDs: {ids}. Cleanup error: {cleanup}') from exc
            raise RuntimeError('Storage failed. New passages were cleaned up; the previous document was preserved.') from exc
        if old_ids:
            try:
                client.delete(collection, points_selector=PointIdsList(points=old_ids), wait=True)
            except Exception as exc:
                raise RuntimeError('New passages were saved, but removing the old document failed. Replacement is incomplete; duplicate passages may remain.') from exc
        yield {'type': 'done', 'filename': name, 'passages': len(chunks), 'replaced': bool(old_ids)}
