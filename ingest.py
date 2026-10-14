import csv  # Standard library module for reading CSV files
import hashlib  # Standard library module for fingerprinting file contents
import uuid  # Standard library module for making unique IDs
from datetime import date, datetime  # Date types, so Excel dates can be written neatly
from pathlib import Path  # Standard library module for working with files and folders

from langchain_core.documents import Document  # LangChain's container for a piece of text plus its metadata
from langchain_text_splitters import RecursiveCharacterTextSplitter  # Cuts long text into smaller chunks
from openpyxl import load_workbook  # Reads Excel (.xlsx) files
from openpyxl.utils import get_column_letter  # Turns a column number into a letter, e.g. 17 -> "Q"
from pypdf import PdfReader  # Reads text out of PDF files

from db import COLLECTION, count_source, delete_source, get_vector_store, stored_sources  # Our helpers for the Qdrant database

DOCS_DIR = Path("documents")  # Folder with the files to add to the database
EXTENSIONS = {".txt", ".md", ".pdf", ".xlsx", ".csv"}  # File types that will be read
TABLE_EXTENSIONS = {".xlsx", ".csv"}  # File types that are tables (read row by row)
PIPELINE_VERSION = "2"  # Bump this when the way files are read changes, so every file is re-added automatically
BATCH_SIZE = 32  # How many chunks to embed at a time (progress is shown after each batch)


def cell_text(value):  # Turns one cell into clean text
    if value is None:  # Empty cell
        return ""  # Nothing
    if isinstance(value, datetime) and value.time() == datetime.min.time():  # A date with no time of day
        return value.date().isoformat()  # e.g. "2024-03-01" instead of "2024-03-01 00:00:00"
    if isinstance(value, (datetime, date)):  # A date with a time
        return value.isoformat(sep=" ")  # e.g. "2024-03-01 14:30:00"
    if isinstance(value, float) and value.is_integer():  # A whole number stored as 12.0
        return str(int(value))  # Write it as 12
    return " ".join(str(value).split())  # Anything else, with extra spaces and line breaks squeezed out


def make_headers(raw):  # Builds clear, unique column names from the header row
    names = [h or f"Column {get_column_letter(i + 1)}" for i, h in enumerate(raw)]  # Blank header -> "Column Q"
    headers = []  # Final names
    for i, name in enumerate(names):  # Go through each column
        if names.count(name) > 1:  # The name appears more than once (e.g. several "Status" columns)
            left = names[i - 1] if i else ""  # The column just to the left
            name = f"{name} ({left})" if left else f"{name} {i + 1}"  # e.g. "Status (2024-03-01)"
        while name in headers:  # Still not unique
            name += f" {i + 1}"  # Add the column number
        headers.append(name)  # Keep it
    return headers  # Hand back the names


def table_rows(grid):  # Cleans a table: returns the column names and the data rows (with their row numbers)
    rows = [(n, [cell_text(v) for v in row]) for n, row in enumerate(grid, start=1)]  # Clean every cell, remembering the row number
    rows = [(n, cells) for n, cells in rows if any(cells)]  # Skip completely empty rows
    if not rows:  # The table has no data
        return [], []  # No columns, no rows
    return make_headers(rows[0][1]), rows[1:]  # First non-empty row is the header; the rest is data


def read_tables(path):  # Reads a spreadsheet file into (sheet name, grid) pairs
    if path.suffix.lower() == ".csv":  # CSV file: one table, no sheet name
        with path.open(encoding="utf-8-sig", errors="ignore", newline="") as f:  # Open it; utf-8-sig handles files saved by Excel
            return [(None, list(csv.reader(f)))]  # The whole file
    workbook = load_workbook(path, data_only=True)  # Excel; data_only gives formula results instead of formulas
    tables = [(ws.title, sheet_grid(ws)) for ws in workbook.worksheets]  # Every sheet
    workbook.close()  # Release the file
    return tables  # Hand back the sheets


def table_documents(grid, path, sheet=None):  # Turns a table into one Document per row
    headers, rows = table_rows(grid)  # Column names and data rows
    where = f"File: {path.name}" + (f" | Sheet: {sheet}" if sheet else "")  # Context line repeated in every row
    docs = []  # One Document per data row
    for n, cells in rows:  # Every data row
        pairs = [f"{h}: {v}" for h, v in zip(headers, cells) if v]  # "Column: value" for each filled cell
        text = f"{where} | Row {n}\n" + " | ".join(pairs)  # Each row is self-contained, so it makes sense on its own
        metadata = {"source": str(path), "row": n}  # Remember the file and row
        if sheet:  # Excel files also have a sheet name
            metadata["sheet"] = sheet  # Remember it
        docs.append(Document(page_content=text, metadata=metadata))  # Add the row
    return docs  # Hand back the rows


def sheet_grid(ws):  # Reads a worksheet into a list of rows, filling merged cells
    grid = [list(row) for row in ws.iter_rows(values_only=True)]  # Every cell value
    for merged in ws.merged_cells.ranges:  # Merged cells only store their value in the top-left cell
        value = grid[merged.min_row - 1][merged.min_col - 1]  # That top-left value
        for r in range(merged.min_row - 1, merged.max_row):  # Every row the merge covers
            for c in range(merged.min_col - 1, merged.max_col):  # Every column the merge covers
                grid[r][c] = value  # Copy the value into it
    return grid  # Hand back the table


def load_file(path):  # Reads one file and returns it as a list of Documents
    suffix = path.suffix.lower()  # File extension, e.g. ".pdf"
    if suffix in TABLE_EXTENSIONS:  # Excel or CSV
        return [d for sheet, grid in read_tables(path) for d in table_documents(grid, path, sheet)]  # One Document per row of every sheet
    if suffix == ".pdf":  # PDFs need a special reader
        reader = PdfReader(path)  # Open the PDF
        return [  # One Document per page, so answers can point to the page
            Document(page_content=page.extract_text() or "", metadata={"source": str(path), "page": n})  # Page text plus file name and page number
            for n, page in enumerate(reader.pages, start=1)  # Go through pages, numbering from 1
        ]  # End of page list
    text = path.read_text(encoding="utf-8", errors="ignore")  # Plain text files are read directly
    return [Document(page_content=text, metadata={"source": str(path)})]  # Wrap it, remembering which file it came from


def find_files():  # Lists every supported file in the documents folder, including subfolders
    return sorted(  # Sorted so the order is always the same
        p for p in DOCS_DIR.rglob("*")  # Look at everything in the documents folder
        if p.suffix.lower() in EXTENSIONS and not p.name.startswith("~$")  # Supported type, and not an Excel/Word temporary lock file
    )  # End of file list


def file_hash(path):  # Fingerprint of a file's contents (changes whenever the file changes)
    return hashlib.sha256(path.read_bytes() + PIPELINE_VERSION.encode()).hexdigest()[:16]  # Short fingerprint


def make_chunks(path):  # Reads one file and cuts it into chunks ready to store
    splitter = RecursiveCharacterTextSplitter(chunk_size=1200, chunk_overlap=200)  # For long text: ~1200 characters, overlapping so ideas aren't cut in half
    chunks = []  # Will hold every chunk
    for doc in load_file(path):  # Each page, row or whole text file
        if "row" in doc.metadata:  # Table rows are already small and complete
            chunks.append(doc)  # Never cut a row in half
        else:  # PDF pages and text files
            chunks.extend(splitter.split_documents([doc]))  # Cut into pieces
    fingerprint = file_hash(path)  # This version of the file
    for c in chunks:  # Label every chunk
        c.metadata["file_hash"] = fingerprint  # So we can tell later whether it's up to date
    return [c for c in chunks if c.page_content.strip()]  # Drop empty chunks (e.g. scanned PDF pages)


def ingest_files(store, files, log=print, force=False):  # Adds files to Qdrant; returns how many chunks were stored
    total = 0  # Chunks stored in this run
    for path in files:  # Go through each file
        chunks = make_chunks(path)  # Read and cut it
        source, fingerprint = str(path), file_hash(path)  # Its name and version
        if not force and chunks and count_source(store, source, fingerprint) == len(chunks):  # Already fully stored and unchanged
            log(f"Up to date: {path} ({len(chunks)} chunks)")  # Nothing to do
            continue  # Next file
        delete_source(store, source)  # Remove old or half-finished chunks of this file
        if not chunks:  # No readable text (e.g. a scanned PDF)
            log(f"Skipped (no text found): {path}")  # Tell the user
            continue  # Next file
        ids = [str(uuid.uuid5(uuid.NAMESPACE_URL, f"{source}:{fingerprint}:{i}")) for i in range(len(chunks))]  # Stable ID per chunk
        log(f"Adding {path}: {len(chunks)} chunks")  # Start message
        for start in range(0, len(chunks), BATCH_SIZE):  # One batch at a time
            end = min(start + BATCH_SIZE, len(chunks))  # Last chunk of this batch
            store.add_documents(chunks[start:end], ids=ids[start:end])  # Turn chunks into vectors and save them
            log(f"  {end}/{len(chunks)} ({end * 100 // len(chunks)}%)")  # Progress
        total += len(chunks)  # Count them
    return total  # Report how many chunks were stored


def remove_missing(store, files, log=print):  # Removes chunks of files that were deleted from the documents folder
    current = {str(p) for p in files}  # Files that exist now
    for source in sorted(stored_sources(store) - current):  # Files in the database but not on disk
        delete_source(store, source)  # Remove their chunks
        log(f"Removed (file no longer exists): {source}")  # Tell the user


def main():  # Entry point of the program
    files = find_files()  # Find everything to add
    store = get_vector_store()  # Open the database
    remove_missing(store, files)  # Clean up deleted files
    if not files:  # Nothing to add
        print(f"No supported files ({', '.join(sorted(EXTENSIONS))}) found in {DOCS_DIR}/")  # Tell the user
        return  # Stop here
    count = ingest_files(store, files)  # Add new and changed files
    print(f"Done. {count} chunks added; the database now holds {store.client.count(COLLECTION).count} chunks.")  # Report


if __name__ == "__main__":  # Only run when this file is executed directly (not imported)
    main()  # Start ingesting
