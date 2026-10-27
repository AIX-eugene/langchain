import io
import threading
import unittest
from unittest.mock import patch

from qdrant_client import QdrantClient
from document_ingestion import extract_documents, ingest_document, read_upload, source_name
from document_ingestion import UPLOAD_COLLECTION as COLLECTION


class Embeddings:
    def embed_documents(self, texts):
        return [[1.0, 0.0, 0.0] for _ in texts]


class IngestionTests(unittest.TestCase):
    def setUp(self):
        self.client = QdrantClient(':memory:')

    def tearDown(self):
        self.client.close()

    def ingest(self, data=b'Guidelines for advertising', name='rules.txt', **kwargs):
        return list(ingest_document(data, name, client=self.client, embeddings=Embeddings(), collection=COLLECTION, **kwargs))

    def test_text_and_metadata(self):
        for name in ['rules.txt', 'rules.md']:
            self.assertEqual(extract_documents(b'Hello', name)[0].metadata, {'source': name})
        self.assertEqual(source_name('C:\\fakepath\\rules.txt'), 'rules.txt')
        with self.assertRaises(ValueError):
            source_name('rules.exe')
        with self.assertRaises(ValueError):
            extract_documents(b'\xff', 'rules.txt')
        with self.assertRaises(ValueError):
            extract_documents(b'  ', 'rules.md')
        with self.assertRaises(ValueError):
            read_upload(io.BytesIO(b''))
        with patch('document_ingestion.MAX_DOCUMENT_MB', 1):
            with self.assertRaises(ValueError):
                read_upload(io.BytesIO(b'a' * (1024*1024+1)))

    def test_office(self):
        from docx import Document
        from openpyxl import Workbook
        doc=Document();doc.add_paragraph('Paragraph');doc.add_table(rows=1, cols=1).cell(0,0).text='Table value'
        buffer=io.BytesIO();doc.save(buffer)
        text=extract_documents(buffer.getvalue(), 'rules.docx')[0].page_content
        self.assertIn('Paragraph',text);self.assertIn('Table value',text)
        book=Workbook();book.active.append(['Item', 0, False]);book.create_sheet('Second').append(['Other'])
        buffer=io.BytesIO();book.save(buffer);book.close()
        docs=extract_documents(buffer.getvalue(),'rules.xlsx')
        self.assertEqual([d.metadata['sheet'] for d in docs], ['Sheet','Second'])
        self.assertEqual(docs[0].metadata['row'],1)
        self.assertIn('False',docs[0].page_content)

    def test_parser_errors_are_actionable(self):
        with patch('openpyxl.load_workbook', side_effect=ImportError('missing dependency')):
            with self.assertLogs('document_ingestion', level='ERROR'):
                with self.assertRaisesRegex(ValueError, 'uv sync'):
                    extract_documents(b'file', 'rules.xlsx')
        with patch('openpyxl.load_workbook', side_effect=RuntimeError('Unsupported workbook structure')):
            with self.assertLogs('document_ingestion', level='ERROR'):
                with self.assertRaisesRegex(ValueError, 'RuntimeError: Unsupported workbook structure'):
                    extract_documents(b'file', 'rules.xlsx')

    def test_pdf(self):
        from fpdf import FPDF
        pdf=FPDF();pdf.add_page();pdf.set_font('Helvetica',size=12);pdf.cell(text='Advertising guidelines')
        self.assertEqual(extract_documents(bytes(pdf.output()),'rules.pdf')[0].metadata['page'],1)
        for name in ['rules.pdf','rules.docx','rules.xlsx']:
            with self.assertRaises(ValueError):
                extract_documents(b'broken file',name)
        from pypdf import PdfWriter
        writer=PdfWriter();writer.add_blank_page(width=100,height=100);writer.encrypt('secret');buffer=io.BytesIO();writer.write(buffer)
        with self.assertRaisesRegex(ValueError,'requires a password'):
            extract_documents(buffer.getvalue(),'rules.pdf')

    def test_aes_pdf_without_open_password(self):
        from fpdf import FPDF
        from pypdf import PdfReader, PdfWriter
        pdf = FPDF(); pdf.add_page(); pdf.set_font('Helvetica', size=12)
        pdf.cell(text='MACEA advertising guidelines')
        writer = PdfWriter()
        writer.append(PdfReader(io.BytesIO(bytes(pdf.output()))))
        writer.encrypt(user_password='', owner_password='owner-secret', algorithm='AES-256')
        buffer = io.BytesIO(); writer.write(buffer)
        chunks = extract_documents(buffer.getvalue(), 'macea.pdf')
        self.assertIn('MACEA advertising guidelines', chunks[0].page_content)
        self.assertEqual(chunks[0].metadata, {'source': 'macea.pdf', 'page': 1})

    def test_duplicate_replace_and_search(self):
        result=self.ingest();self.assertEqual(result[-1]['type'],'done')
        self.assertEqual(self.ingest()[-1]['type'],'duplicate')
        self.assertTrue(self.ingest(b'Updated guidelines',replace=True)[-1]['replaced'])
        points=self.client.query_points(COLLECTION,query=[1.,0.,0.],limit=3).points
        self.assertEqual(len(points),1)
        self.assertEqual(points[0].payload['content'],'Updated guidelines')

    def test_failed_storage_preserves_old(self):
        self.ingest()
        original=self.client.upsert
        def fail(*args,**kwargs):
            original(*args,**kwargs)
            raise RuntimeError('connection lost after write')
        with patch.object(self.client,'upsert',side_effect=fail):
            with self.assertRaisesRegex(RuntimeError,'cleaned up'):
                self.ingest(b'New rules',replace=True)
        self.assertEqual(self.client.count(COLLECTION).count,1)
        self.assertEqual(self.client.scroll(COLLECTION)[0][0].payload['content'],'Guidelines for advertising')

    def test_failed_embedding_preserves_old(self):
        self.ingest()
        with patch.object(Embeddings,'embed_documents',side_effect=RuntimeError('Ollama unavailable')):
            with self.assertRaisesRegex(RuntimeError,'Ollama unavailable'):
                self.ingest(b'New rules',replace=True)
        self.assertEqual(self.client.count(COLLECTION).count,1)

    def test_incomplete_replacement_reported(self):
        self.ingest()
        with patch.object(self.client,'delete',side_effect=RuntimeError('unavailable')):
            with self.assertRaisesRegex(RuntimeError,'Replacement is incomplete'):
                self.ingest(b'New rules',replace=True)
        self.assertEqual(self.client.count(COLLECTION).count,2)

    def test_incomplete_cleanup_reported(self):
        self.ingest()
        with patch.object(self.client, 'upsert', side_effect=RuntimeError('write failed')):
            with patch.object(self.client, 'delete', side_effect=RuntimeError('cleanup failed')):
                with self.assertRaisesRegex(RuntimeError, 'cleanup is incomplete'):
                    self.ingest(b'New rules', replace=True)

    def test_multiple_batches_and_textless_pdf(self):
        data = ('Advertising rules apply. ' * 2000).encode()
        result = self.ingest(data)
        self.assertGreater(result[-1]['passages'], 32)
        self.assertEqual(self.client.count(COLLECTION).count, result[-1]['passages'])
        from pypdf import PdfWriter
        writer = PdfWriter(); writer.add_blank_page(width=100, height=100)
        buffer = io.BytesIO(); writer.write(buffer)
        with self.assertRaisesRegex(ValueError, 'OCR'):
            extract_documents(buffer.getvalue(), 'scan.pdf')

    def test_concurrent_uploads(self):
        results=[]
        # Local Qdrant supports this when SQLite thread checks are disabled.
        client=QdrantClient(':memory:',force_disable_check_same_thread=True)
        def run():
            results.append(list(ingest_document(b'Rules','same.txt',client=client,embeddings=Embeddings(),collection=COLLECTION))[-1]['type'])
        threads=[threading.Thread(target=run) for _ in range(2)]
        for thread in threads:thread.start()
        for thread in threads:thread.join()
        self.assertCountEqual(results,['done','duplicate']);client.close()


if __name__=='__main__':
    unittest.main()
